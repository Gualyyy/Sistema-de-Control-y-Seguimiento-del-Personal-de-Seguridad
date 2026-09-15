from docx import Document
from docx.shared import Pt

# Crear documento
doc = Document()

# Título
title = doc.add_heading('Documentación del Sistema: Inicio a Fin', level=1)

# Resumen
doc.add_heading('Resumen', level=2)
doc.add_paragraph(
    'Documento que describe el flujo completo del sistema de autenticación, registro y panel de usuario.'
)

# Índice
doc.add_heading('Índice', level=2)
doc.add_paragraph('1. Visión general')
doc.add_paragraph('2. Estructura del proyecto')
doc.add_paragraph('3. Flujo de autenticación')
doc.add_paragraph('4. Protección de rutas y tokenización del panel')
doc.add_paragraph('5. Gestión de sesiones y cierre')
doc.add_paragraph('6. Formularios y validaciones')
doc.add_paragraph('7. Archivos importantes (plantillas, estáticos, vistas)')
doc.add_paragraph('8. Pruebas automatizadas')
doc.add_paragraph('9. Despliegue y notas operativas')

# Secciones

doc.add_heading('1. Visión general', level=2)
doc.add_paragraph(
    'El sistema es una aplicación Django llamada "vigilancia" que centraliza autenticación, gestión de usuarios y ' \
    'asignación de equipos. Protege las rutas sensibles mediante sesiones y un token firmado para el panel de usuario, ' \
    'ofrece formularios para completar datos y utiliza validaciones tanto en cliente (JS) como en servidor (views).' 
)

doc.add_heading('2. Estructura del proyecto', level=2)
paragraph = doc.add_paragraph()
paragraph.add_run('Estructura principal:').bold = True
doc.add_paragraph(' - manage.py')
doc.add_paragraph(' - inicio_sesion/ (app principal con vistas, templates, static, models, tests)')
doc.add_paragraph(' - vigilancia/ (configuración del proyecto Django)')

doc.add_heading('3. Flujo de autenticación', level=2)
doc.add_paragraph('1) El usuario accede a la página de login (ruta raíz).')
doc.add_paragraph('2) Credenciales se envían vía POST a `inicio_sesion` (views/autenticacion.py).')
doc.add_paragraph('3) Si la autenticación es correcta:')
para = doc.add_paragraph('   - Se crean variables de sesión: `usuario`, `perfil`, `cedula_identidad` y `panel_token` (token firmado).')
para = doc.add_paragraph('   - Si el registro de nacimiento está incompleto, la respuesta retorna `estado: incompleto` con URL `completar_registro`.')
para = doc.add_paragraph('   - Se registra la sesión en la tabla `sesiones_usuario` (se crea o actualiza el registro existente para evitar claves duplicadas).')

doc.add_heading('4. Protección de rutas y tokenización del panel', level=2)
doc.add_paragraph(' - Decorador `autenticado_required` (inicio_sesion/views/autenticacion.py) comprueba: existencia de `usuario` y `panel_token` en la sesión, y valida token con `validar_token_panel`.')
doc.add_paragraph(' - La vista `panel_usuario` genera (o reutiliza) el token y redirige a `panel/<token>/` (panel_usuario_token).')
doc.add_paragraph(' - `panel_usuario_token` verifica que el token de la URL coincida con el de la sesión y aplica cabeceras `no-store` y `no-cache` para evitar cacheo.')

doc.add_heading('5. Gestión de sesiones y cierre', level=2)
doc.add_paragraph(' - Cuando el usuario inicia sesión se registra la sesión en `sesiones_usuario`. Si existe un registro con la misma `session_key`, se actualiza para evitar `IntegrityError` por clave duplicada.')
doc.add_paragraph(' - Al cerrar sesión via `cerrar_sesion`, se marca la sesión como no activa (`activa=False`), se pone `fecha_fin`, y se limpian todas las claves de sesión (se llama `request.session.flush()`).')

doc.add_heading('6. Formularios y validaciones', level=2)
doc.add_paragraph(' - Formulario "Completar Datos Empleado": validaciones en cliente (JS) y servidor:')
doc.add_paragraph('   - Teléfonos: exactamente 7 dígitos')
doc.add_paragraph('   - Correos: el usuario solo permite letras, números y un punto; se selecciona dominio en un select')
doc.add_paragraph('   - Fecha de nacimiento: validación para asegurar edad mínima de 18 años (cliente y servidor).')
doc.add_paragraph('   - Mensajes cliente: se usan SweetAlert2 para mostrar alertas consistentes.')

doc.add_heading('7. Archivos importantes', level=2)
doc.add_paragraph(' - Vistas: inicio_sesion/views/autenticacion.py, empleados.py, panel.py')
doc.add_paragraph(' - URLs: inicio_sesion/urls.py (rutas protegidas con `autenticado_required`)')
doc.add_paragraph(' - Templates clave: inicio_sesion/Templates/inicio_sesion/completar_datos_usuario.html, autenticacion.html, panel_usuario.html')
doc.add_paragraph(' - JS cliente: static/Funcionalidades (validación de correo, teléfono, residencia, nacimiento) y Funcionalidades_Asincronas (completar_registrar.js, autenticacion.js)')

doc.add_heading('8. Pruebas automatizadas', level=2)
doc.add_paragraph(' - tests en inicio_sesion/tests.py cubren: protección de rutas, redirección al panel tokenizado, validación de edad, y manejo de sesiones duplicadas.')

doc.add_heading('9. Despliegue y notas operativas', level=2)
doc.add_paragraph(' - Asegurar configuración de base de datos y SECRET_KEY en producción. Use HTTPS y configuraciones seguras de sesión (SESSION_COOKIE_SECURE, CSRF settings).')
doc.add_paragraph(' - Mantener rotación de tokens si fuera necesario y considerar invalidación de sesiones global en caso de eventos de seguridad.')

# Footer
doc.add_paragraph('\nGenerado automáticamente: Documentación de sistema - flujo de autenticación y registro.')

# Guardar documento
output_path = 'Documentacion_Sistema.docx'
doc.save(output_path)
print('Saved', output_path)
